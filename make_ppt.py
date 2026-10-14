import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

R = json.load(open("outputs/results.json"))
M = R["models"]
NAVY, RED, GREY, LIGHT, WHITE = RGBColor(0x14, 0x24, 0x4a), RGBColor(0xd9, 0x53, 0x4f), RGBColor(0x55, 0x5f, 0x70), RGBColor(0xf1, 0xf4, 0xf9), RGBColor(255, 255, 255)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def box(s, x, y, w, h, fill=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.line.fill.background()
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    return sh


def text(s, x, y, w, h, lines, size=18, color=NAVY, bold=False, align=None, bullet=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("•  " if bullet else "") + ln
        p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold; p.space_after = Pt(8)
        if align: p.alignment = align
    return tb


def slide(title, sub=None):
    s = prs.slides.add_slide(blank)
    box(s, 0, 0, 13.333, 1.15, NAVY); box(s, 0, 1.15, 13.333, 0.06, RED)
    text(s, 0.6, 0.18, 12, 0.8, title, 30, WHITE, True)
    if sub: text(s, 0.6, 1.3, 12, 0.5, sub, 16, GREY)
    return s


def pic(s, path, x, y, w=None, h=None):
    s.shapes.add_picture("outputs/" + path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def stat(s, x, y, big, small, w=2.9):
    box(s, x, y, w, 1.5, LIGHT)
    text(s, x, y + 0.1, w, 0.8, big, 34, RED, True, PP_ALIGN.CENTER)
    text(s, x, y + 0.9, w, 0.5, small, 13, GREY, False, PP_ALIGN.CENTER)


def table(s, x, y, w, rows, colw, size=13):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.4 * len(rows))).table
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = str(v); p = c.text_frame.paragraphs[0]; p.font.size = Pt(size)
            p.font.bold = i == 0 or False; p.font.color.rgb = WHITE if i == 0 else NAVY
            c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else WHITE)


# 1 title
s = prs.slides.add_slide(blank); box(s, 0, 0, 13.333, 7.5, NAVY); box(s, 0.8, 3.55, 2.2, 0.07, RED)
text(s, 0.8, 1.8, 11.5, 1.6, "Disaster Tweet Classification", 48, WHITE, True)
text(s, 0.8, 3.8, 11.5, 1, "An NLP project: from raw tweets to a deployed web classifier", 22, RGBColor(0xc9, 0xd3, 0xe8))
text(s, 0.8, 6.4, 11.5, 0.5, "Digicrome AI Course · Submission 15 Oct 2026", 14, RGBColor(0xc9, 0xd3, 0xe8))

# 2 problem
s = slide("The problem", "Telling real emergencies from metaphors on social media")
text(s, 0.6, 1.9, 6.3, 4, ["Relief agencies and newsrooms want automated Twitter monitoring for disasters.",
     "Keywords alone fail: \"this song is fire\" is not a wildfire.",
     "Goal: high precision (few false alarms) and good recall, robust to slang and noise, fast enough to scale."], 19, bullet=True)
box(s, 7.4, 1.9, 5.3, 3.9, LIGHT)
text(s, 7.7, 2.05, 4.8, 0.5, "Same word, different intent", 18, RED, True)
text(s, 7.7, 2.7, 4.8, 3, ["Forest fire near La Ronge Sask. Canada  →  disaster", "On top of the hill, I can see a fire in the woods  →  disaster",
     "My car is so fast  →  not a disaster", "I'm flooded with work  →  not a disaster"], 15, NAVY)

# 3 approach
s = slide("Workflow at a glance", "Four parts of the brief, one reproducible pipeline")
steps = [("1", "Explore & prepare", "EDA, cleaning, tokenising, split"), ("2", "Features & models", "TF-IDF, extras, 5 models, grid search"),
         ("3", "Evaluate & validate", "Metrics, ROC/PR, overfit & robustness"), ("4", "Deploy", "Pickle + Flask web UI")]
for i, (n, h, d) in enumerate(steps):
    x = 0.6 + i * 3.1
    box(s, x, 2.3, 2.85, 3.2, LIGHT); box(s, x, 2.3, 2.85, 0.08, RED)
    text(s, x + 0.2, 2.5, 2.4, 0.8, n, 40, RED, True); text(s, x + 0.2, 3.3, 2.5, 0.6, h, 19, NAVY, True); text(s, x + 0.2, 4.1, 2.5, 1.3, d, 15, GREY)
text(s, 0.6, 6.0, 12, 0.6, "Code: pipeline.py (analysis & training) · app/ (Flask) · make_ppt.py (this deck)", 14, GREY)

# 4 data
s = slide("Data exploration", "tweets.csv: id, keyword, location, text, target")
stat(s, 0.6, 1.9, f"{R['shape'][0]:,}", "labelled tweets (brief says ~10k)")
stat(s, 0.6, 3.6, f"{R['class_counts']['1']/R['shape'][0]:.0%}", "disaster (target = 1)")
stat(s, 0.6, 5.3, f"{R['dups']}", f"duplicate texts; {R['conflicting_dup_texts']} with conflicting labels", 2.9)
pic(s, "class_dist.png", 3.8, 1.8, h=3.9)
text(s, 8.3, 1.9, 4.6, 4.5, [f"Keyword missing in {R['nulls']['keyword']} rows, location in {R['nulls']['location']} (34%), so location is dropped.",
     "Mild imbalance (43/57), so F1 and PR curves are tracked, not just accuracy.",
     "Exact duplicate texts removed (keep first) to prevent train/test leakage: 7,503 rows remain."], 15, bullet=True)

# 5 keywords
s = slide("Which words signal a disaster?", "Top non-stop-words per class after cleaning")
pic(s, "top_words.png", 0.5, 1.8, w=8.2)
text(s, 8.9, 1.9, 4.1, 4.8, ["Disaster: killed, california, suicide (bombers), police, fire, flood, disaster, news, storm.",
     "Non-disaster: everyday chatter (like, love, new, just).",
     "Many words (fire, burning, bomb) appear in both classes, which is why context features matter.",
     ], 15, bullet=True)

# 6 preparation
s = slide("Data preparation")
text(s, 0.6, 1.6, 6.2, 5, ["Lowercase; strip URLs, @mentions, '#', HTML entities, digits, punctuation.", "Tokenise into words (len ≥ 2), used for frequency analysis and extra features.",
     "Labels already numeric: 0 = non-disaster, 1 = disaster.", f"Stratified 80/20 split (random_state=42): {R['train']:,} train / {R['test']:,} test.",
     "All vectorisers, scalers and SVD are fit on train only, so there is no leakage."], 18, bullet=True)
box(s, 7.2, 1.7, 5.5, 4.3, LIGHT)
text(s, 7.4, 1.8, 5.1, 0.4, "Before → after", 16, RED, True)
text(s, 7.4, 2.3, 5.1, 3.6, ["#RockyFire Update => California Hwy. 20 closed in both directions due to Lake County fire - #CAfire #wildfires",
     "↓", "rockyfire update california hwy closed in both directions due to lake county fire cafire wildfires"], 15, NAVY)

# 7 features
s = slide("Feature engineering", "5-fold CV F1 of logistic regression on the training set")
ab = R["feature_ablation_f1_cv"]
table(s, 0.6, 1.9, 7.0, [["Feature set", "CV F1"], ["Tweet-statistics only (length, #, @, URL, !, ?, caps, sentiment)", f"{R['extras_only_f1_cv']:.3f}"],
      ["LSA vectors (150-d SVD, semantic embedding stand-in)", f"{R['lsa_only_f1_cv']:.3f}"], ["TF-IDF words (1-2 grams) + keyword", f"{ab['TF-IDF words']:.3f}"],
      ["+ character 2-5 grams (slang, typos)", f"{ab['TF-IDF word+char']:.3f}"], ["+ tweet statistics (chosen)", f"{ab['TF-IDF word+char+extra']:.3f}"]], [5.6, 1.4])
text(s, 8.0, 1.9, 4.8, 4.6, ["TF-IDF is the strongest single signal; char n-grams add robustness to slang and misspellings.",
     "Tweet statistics help only slightly alone, but add a small gain on top.",
     "Word2Vec/GloVe need extra downloads (no gensim installed), so LSA is the semantic-vector substitute here. Pretrained GloVe is the obvious next step.",
     "Sentiment is a small lexicon score (positive minus negative words)."], 14, bullet=True)

# 8 models
s = slide("Model selection and tuning", "Five candidates, 5-fold CV grid search optimising F1")
rows = [["Model", "Best params", "CV F1", "Train F1"]]
for n, m in M.items(): rows.append([n, ", ".join(f"{k}={v}" for k, v in m["best_params"].items()), f"{m['cv_f1']:.3f}", f"{m['train_f1']:.3f}"])
table(s, 0.6, 1.9, 12.1, rows, [3.8, 4.7, 1.8, 1.8], 14)
text(s, 0.6, 4.7, 12, 1.8, ["Linear models on sparse TF-IDF lead (CV F1 ≈ 0.76); Random Forest and the MLP on dense LSA are weaker on short, sparse text.",
     "Train F1 far above CV F1 for LR, SVM and RF shows overfitting, addressed in validation."], 16, bullet=True)

# 9 evaluation
s = slide("Model evaluation on the held-out test set", f"{R['test']:,} unseen tweets")
rows = [["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"]]
for n, m in M.items(): rows.append([n, *(f"{m[k]:.3f}" for k in ["accuracy", "precision", "recall", "f1", "roc_auc"])])
table(s, 0.6, 1.9, 7.4, rows, [2.8, 0.9, 0.9, 0.9, 0.9, 1.0], 12)
pic(s, "model_compare.png", 8.2, 1.8, w=4.8)
text(s, 0.6, 4.7, 7.4, 2, [f"Winner: {R['best_model']} (F1 {M[R['best_model']]['f1']:.3f}, accuracy {M[R['best_model']]['accuracy']:.3f}).",
     "Random Forest has the highest precision (0.876) but misses 38% of disasters; choose by cost of false alarms vs misses."], 15, bullet=True)

# 10 curves
s = slide("ROC, precision-recall and confusion matrices")
pic(s, "roc_pr.png", 0.5, 1.4, w=8.2)
pic(s, "confusion_best.png", 8.9, 1.4, w=4.0)
bc = R["best_confusion"]
text(s, 8.9, 4.5, 4.1, 2.6, [f"Best model: {bc[1][1]} true alerts, {bc[0][1]} false alarms, {bc[1][0]} missed disasters, {bc[0][0]} correct rejections."], 14)
pic(s, "confusion_all.png", 0.5, 5.3, w=8.2)

# 11 validation
s = slide("Model validation", "Generalisation, over/underfitting and robustness")
pic(s, "learning_curve.png", 0.5, 1.9, h=3.5)
text(s, 6.4, 1.6, 6.5, 5, [f"Test F1 {R['deploy_model_test']['f1']:.3f} matches CV F1 {M['Logistic Regression']['cv_f1']:.3f}: the estimate is stable on unseen data.",
     f"Overfitting is real: train F1 {R['learning_curve']['train']:.2f} vs validation {R['learning_curve']['cv']:.2f}. Regularisation (C) was tuned; the CV curve still rises with more data, so more labelled tweets is the best fix.",
     f"Noise test (caps, repeated letters, junk URL/hashtag/mention added): F1 {R['deploy_model_test']['f1']:.3f} → {R['robust_f1_noisy']:.3f}, a drop of only {R['deploy_model_test']['f1']-R['robust_f1_noisy']:.3f}.",
     "Not underfitting: it beats the extras-only and LSA baselines clearly."], 15, bullet=True)
text(s, 0.6, 5.7, 5.6, 1.2, "Caveat: labels are noisy (some errors are arguably mislabelled), which caps achievable accuracy near the low 80s.", 13, GREY)

# 12 interpretation
s = slide("What the model learned, and where it fails")
text(s, 0.6, 1.5, 4, 0.4, "Strongest disaster words", 17, RED, True)
text(s, 0.6, 2.0, 4, 3, [f"{w}  ({c:+.1f})" for w, c in R["top_coef_pos"][:8]], 14)
text(s, 4.6, 1.5, 4, 0.4, "Strongest non-disaster words", 17, RED, True)
text(s, 4.6, 2.0, 4, 3, [f"{w}  ({c:+.1f})" for w, c in R["top_coef_neg"][:8]], 14)
text(s, 8.6, 1.5, 4.3, 0.4, "Metaphor probes (P disaster)", 17, RED, True)
text(s, 8.6, 2.0, 4.5, 4.5, [f"{p}  →  {v:.2f}" for p, v in R["probes"]], 12)
text(s, 0.6, 5.2, 7.8, 1.8, ["Errors cluster in political/war talk, sports-riot comments, and promotional spam that reuses disaster vocabulary.",
     "Figurative uses (\"on fire\", \"bomb\") are still sometimes over-scored; contextual embeddings (BERT) are the planned upgrade."], 14, bullet=True)

# 13 deployment
s = slide("Deployment: model + Flask web app", "Part 4 of the brief")
text(s, 0.6, 1.9, 6.4, 5, ["Final model: one scikit-learn Pipeline (clean → word+char TF-IDF → logistic regression, C=3), pickled with joblib as app/model.pkl. It takes raw text only, so it skips the extra tweet statistics (test F1 0.767 vs 0.775 with them).",
     "Flask backend: POST /predict takes JSON text and returns label and probability; GET / serves the UI.",
     "UI: text box, Classify button, colour-coded verdict with probability bar, one-click examples, empty-input error.",
     "Run locally:  pip install -r app/requirements.txt, then python app/app.py.",
     "Cloud: gunicorn app:app on Heroku, Render or AWS Elastic Beanstalk (Procfile / requirements provided). Not deployed to a live host from here."], 15, bullet=True)
box(s, 7.4, 1.9, 5.3, 3.9, LIGHT)
text(s, 7.7, 2.0, 4.8, 0.4, "Web UI", 17, RED, True)
text(s, 7.7, 2.6, 4.8, 3, ["[ Paste a tweet here…            ]", "[ Classify ]", "", "Real disaster, probability 98.8%", "████████████████░"], 15, NAVY)
text(s, 7.7, 5.0, 4.8, 0.7, "Tested via Flask test client: / → 200, valid text → 98.8%, empty → 400.", 12, GREY)

# 14 scalability + conclusion
s = slide("Scalability, limits and next steps")
text(s, 0.6, 1.6, 6, 0.4, "Scalability", 19, RED, True)
text(s, 0.6, 2.1, 6, 4, ["Inference is a sparse linear dot product: thousands of tweets/sec per core.", "Stateless app: scale horizontally behind gunicorn workers.", "Batch endpoint and streaming ingestion are straightforward additions."], 15, bullet=True)
text(s, 6.9, 1.6, 6, 0.4, "Next steps", 19, RED, True)
text(s, 6.9, 2.1, 6, 4, ["Fine-tune a transformer (DistilBERT) to resolve metaphors.", "Use pretrained GloVe/Word2Vec features.", "Tune the decision threshold for the agency's false-alarm tolerance.", "Add location/time and retweet-spike signals; monitor drift."], 15, bullet=True)
box(s, 0.6, 5.5, 12.1, 1.3, NAVY)
text(s, 0.9, 5.65, 11.5, 1.1, f"Result: an end-to-end classifier with {R['deploy_model_test']['accuracy']:.0%} accuracy, {R['deploy_model_test']['precision']:.0%} precision and F1 {R['deploy_model_test']['f1']:.2f} on unseen tweets (deployed model), deployed behind a working web UI.", 18, WHITE, True)

prs.save("Disaster_Tweet_Classification.pptx")
print("saved")
